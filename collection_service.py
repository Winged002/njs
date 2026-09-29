import csv
import io
import ipaddress
import json
import mimetypes
import os
import re
import socket
import uuid
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup

from config import Config

ALLOWED_DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".odt", ".xls", ".xlsx", ".csv", ".txt", ".md", ".rtf"}


def _now():
    return datetime.now(timezone.utc)


def _plain(value, limit=20000):
    text = BeautifulSoup(str(value or ""), "html.parser").get_text(" ", strip=True)
    text = re.sub(r"\s+", " ", text).strip()
    return text[:limit]


def _validate_public_url(value):
    raw = str(value or "").strip()
    if not raw:
        raise ValueError("Enter a URL")
    parsed = urlparse(raw if "://" in raw else "https://" + raw)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise ValueError("Use a normal http or https URL")
    host = parsed.hostname.lower().rstrip(".")
    if host in {"localhost", "localhost.localdomain"}:
        raise ValueError("Local/private URLs cannot be collected")
    try:
        infos = socket.getaddrinfo(host, parsed.port or (443 if parsed.scheme == "https" else 80), type=socket.SOCK_STREAM)
    except socket.gaierror as exc:
        raise ValueError("The URL hostname could not be resolved") from exc
    for info in infos:
        ip = ipaddress.ip_address(info[4][0])
        if not ip.is_global:
            raise ValueError("Local/private URLs cannot be collected")
    return parsed.geturl()


def _fetch_html(url):
    current = _validate_public_url(url)
    headers = {
        "User-Agent": "NewsjackingCore/1.9 (+collection research)",
        "Accept": "text/html,application/xhtml+xml,text/plain;q=0.8,*/*;q=0.1",
    }
    session = requests.Session()
    response = None
    for _ in range(5):
        response = session.get(
            current,
            headers=headers,
            timeout=Config.WEBSITE_RESEARCH_TIMEOUT_SECONDS,
            allow_redirects=False,
            stream=True,
        )
        if response.status_code in {301, 302, 303, 307, 308}:
            target = response.headers.get("Location")
            response.close()
            if not target:
                raise RuntimeError("URL redirected without a destination")
            current = _validate_public_url(urljoin(current, target))
            continue
        break
    if response is None:
        raise RuntimeError("URL request failed")
    response.raise_for_status()
    content_type = (response.headers.get("content-type") or "").lower()
    if "html" not in content_type and not content_type.startswith("text/"):
        response.close()
        raise ValueError("The URL did not return a web page")
    max_bytes = int(getattr(Config, "COLLECTION_URL_MAX_BYTES", Config.WEBSITE_RESEARCH_MAX_BYTES))
    chunks = []
    total = 0
    for chunk in response.iter_content(65536):
        if not chunk:
            continue
        total += len(chunk)
        if total > max_bytes:
            break
        chunks.append(chunk)
    encoding = response.encoding or "utf-8"
    response.close()
    return current, b"".join(chunks).decode(encoding, errors="replace")


def scrape_url(url):
    final_url, raw = _fetch_html(url)
    soup = BeautifulSoup(raw, "html.parser")
    title = _plain(soup.title.string if soup.title and soup.title.string else final_url, 300)
    meta = soup.find("meta", attrs={"name": re.compile("^description$", re.I)})
    description = _plain(meta.get("content") if meta else "", 1500)
    headings = [_plain(node.get_text(" "), 350) for node in soup.find_all(["h1", "h2", "h3"])[:40]]

    visible = BeautifulSoup(raw, "html.parser")
    for node in visible(["script", "style", "noscript", "template", "svg"]):
        node.decompose()
    body = _plain(visible.get_text(" ", strip=True), int(getattr(Config, "COLLECTION_ITEM_TEXT_MAX_CHARS", 50000)))

    base_host = (urlparse(final_url).hostname or "").lower()
    internal, external, seen = [], [], set()
    max_links = int(getattr(Config, "COLLECTION_DISCOVERED_LINK_LIMIT", 120))
    for anchor in soup.find_all("a", href=True):
        href = urljoin(final_url, anchor.get("href"))
        parsed = urlparse(href)
        if parsed.scheme not in {"http", "https"} or not parsed.hostname:
            continue
        href = href.split("#", 1)[0]
        if href in seen or href == final_url:
            continue
        seen.add(href)
        row = {"url": href, "label": _plain(anchor.get_text(" "), 160), "kind": "internal" if parsed.hostname.lower() == base_host else "external"}
        (internal if row["kind"] == "internal" else external).append(row)
        if len(seen) >= max_links:
            break

    images, image_seen = [], set()
    for node in soup.find_all("img"):
        src = node.get("src") or node.get("data-src") or node.get("data-lazy-src")
        if not src:
            continue
        src = urljoin(final_url, src)
        parsed = urlparse(src)
        if parsed.scheme not in {"http", "https"} or src in image_seen:
            continue
        image_seen.add(src)
        images.append({
            "url": src,
            "alt": _plain(node.get("alt"), 300),
            "width": str(node.get("width") or "")[:20],
            "height": str(node.get("height") or "")[:20],
        })
        if len(images) >= int(getattr(Config, "COLLECTION_IMAGE_LIMIT", 40)):
            break

    return {
        "url": url,
        "final_url": final_url,
        "title": title,
        "description": description,
        "headings": headings,
        "text": body,
        "internal_links": internal,
        "external_links": external,
        "images": images,
        "scraped_at": _now(),
    }


def youtube_video_id(url):
    parsed = urlparse(str(url or "").strip())
    host = (parsed.hostname or "").lower()
    if host in {"youtu.be", "www.youtu.be"}:
        return parsed.path.strip("/").split("/", 1)[0]
    if "youtube.com" in host:
        if parsed.path == "/watch":
            from urllib.parse import parse_qs
            return (parse_qs(parsed.query).get("v") or [""])[0]
        match = re.search(r"/(?:shorts|embed|live)/([^/?#]+)", parsed.path)
        if match:
            return match.group(1)
    return ""


def fetch_supadata_transcript(url):
    key = str(getattr(Config, "SUPADATA_API_KEY", "") or "").strip()
    if not key:
        raise RuntimeError("SUPADATA_API_KEY is not configured")
    base = str(getattr(Config, "SUPADATA_BASE_URL", "https://api.supadata.ai/v1") or "").rstrip("/")
    headers = {"x-api-key": key, "Accept": "application/json"}
    # Current Supadata unified transcript endpoint. It supports YouTube URLs and can return text directly.
    response = requests.get(f"{base}/transcript", params={"url": url, "text": "true"}, headers=headers, timeout=60)
    if response.status_code >= 400:
        # Compatibility fallback for accounts using the dedicated YouTube transcript endpoint.
        vid = youtube_video_id(url)
        if not vid:
            response.raise_for_status()
        response = requests.get(f"{base}/youtube/transcript", params={"videoId": vid, "text": "true", "lang": "en"}, headers=headers, timeout=60)
    response.raise_for_status()
    payload = response.json()
    content = payload.get("content")
    if isinstance(content, list):
        segments = [
            {
                "text": _plain(item.get("text"), 5000),
                "offset": item.get("offset", item.get("start")),
                "duration": item.get("duration"),
            }
            for item in content if isinstance(item, dict) and item.get("text")
        ]
        text = " ".join(segment["text"] for segment in segments)
    else:
        text = _plain(content, int(getattr(Config, "COLLECTION_TRANSCRIPT_MAX_CHARS", 160000)))
        segments = []
    if not text:
        raise RuntimeError("Supadata returned no transcript text")
    return {
        "language": payload.get("lang"),
        "available_languages": payload.get("availableLangs") or [],
        "text": text,
        "segments": segments[:4000],
        "provider": "supadata",
        "fetched_at": _now(),
    }


def _chunk_text(text, target=6000):
    text = re.sub(r"\s+", " ", str(text or "")).strip()
    if not text:
        return []
    chunks = []
    while text:
        if len(text) <= target:
            chunks.append(text)
            break
        cut = text.rfind(" ", 0, target)
        if cut < target // 2:
            cut = target
        chunks.append(text[:cut].strip())
        text = text[cut:].strip()
    return chunks


def extract_document(path, original_name=""):
    path = Path(path)
    ext = path.suffix.lower()
    if ext not in ALLOWED_DOCUMENT_EXTENSIONS:
        raise ValueError(f"Unsupported document type: {ext or 'unknown'}")
    title = original_name or path.name
    items = []

    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        for index, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if _plain(text):
                items.append({"item_type": "document_section", "title": f"{title} · page {index + 1}", "text": _plain(text, 50000), "metadata": {"page": index + 1}})
    elif ext == ".docx":
        from docx import Document
        doc = Document(str(path))
        raw = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        for index, chunk in enumerate(_chunk_text(raw)):
            items.append({"item_type": "document_section", "title": f"{title} · section {index + 1}", "text": chunk, "metadata": {"section": index + 1}})
    elif ext == ".odt":
        from odf.opendocument import load
        from odf import text as odf_text
        from odf.teletype import extractText
        doc = load(str(path))
        raw = "\n".join(extractText(paragraph) for paragraph in doc.getElementsByType(odf_text.P))
        for index, chunk in enumerate(_chunk_text(raw)):
            items.append({"item_type": "document_section", "title": f"{title} · section {index + 1}", "text": chunk, "metadata": {"section": index + 1}})
    elif ext == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(str(path), read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                rows.append(["" if cell is None else str(cell) for cell in row])
                if len(rows) >= 1000:
                    break
            out = io.StringIO()
            writer = csv.writer(out)
            writer.writerows(rows)
            items.append({"item_type": "table", "title": f"{title} · {sheet.title}", "text": out.getvalue()[:100000], "metadata": {"sheet": sheet.title, "rows": len(rows)}})
    elif ext == ".xls":
        import xlrd
        workbook = xlrd.open_workbook(str(path), on_demand=True)
        for sheet_name in workbook.sheet_names():
            sheet = workbook.sheet_by_name(sheet_name)
            rows = [[str(sheet.cell_value(r, c)) for c in range(sheet.ncols)] for r in range(min(sheet.nrows, 1000))]
            out = io.StringIO(); csv.writer(out).writerows(rows)
            items.append({"item_type": "table", "title": f"{title} · {sheet_name}", "text": out.getvalue()[:100000], "metadata": {"sheet": sheet_name, "rows": len(rows)}})
    elif ext == ".csv":
        raw = path.read_text(encoding="utf-8", errors="replace")
        items.append({"item_type": "table", "title": title, "text": raw[:120000], "metadata": {"format": "csv"}})
    else:
        raw = path.read_text(encoding="utf-8", errors="replace")
        for index, chunk in enumerate(_chunk_text(raw)):
            items.append({"item_type": "document_section", "title": f"{title} · section {index + 1}", "text": chunk, "metadata": {"section": index + 1}})

    if not items:
        items.append({"item_type": "document_section", "title": title, "text": "", "metadata": {"warning": "No extractable text found"}})
    return items


def item_prompt_view(item):
    metadata = dict(item.get("metadata") or {})
    if "segments" in metadata:
        metadata["segment_count"] = len(metadata.get("segments") or [])
        metadata.pop("segments", None)
    return {
        "id": str(item.get("_id")),
        "collection_id": str(item.get("collection_id")),
        "source_id": str(item.get("source_id")),
        "type": item.get("item_type"),
        "title": item.get("title"),
        "text": str(item.get("text") or "")[:16000],
        "url": item.get("url"),
        "image_url": item.get("image_url"),
        "metadata": metadata,
    }


def evidence_context(items, max_chars=50000):
    payload = []
    used = 0
    for item in items:
        row = item_prompt_view(item)
        encoded = json.dumps(row, ensure_ascii=False, default=str)
        if used + len(encoded) > max_chars:
            remaining = max_chars - used
            if remaining < 1000:
                break
            row["text"] = str(row.get("text") or "")[: max(0, remaining - 500)]
            encoded = json.dumps(row, ensure_ascii=False, default=str)
        payload.append(row)
        used += len(encoded)
    return payload
